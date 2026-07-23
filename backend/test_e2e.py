"""Uçtan uca test. Backend çalışırken: python test_e2e.py"""

import io
import json
import urllib.error
import urllib.parse
import urllib.request

BASE = "http://localhost:8000/api"


def call(method, path, token=None, data=None, form=None, raw=None, raw_ct=None):
    url = BASE + path
    headers = {}
    body = None
    if form:
        body = urllib.parse.urlencode(form).encode()
        headers["Content-Type"] = "application/x-www-form-urlencoded"
    elif data is not None:
        body = json.dumps(data).encode()
        headers["Content-Type"] = "application/json"
    elif raw is not None:
        body = raw
        headers["Content-Type"] = raw_ct
    if token:
        headers["Authorization"] = f"Bearer {token}"
    req = urllib.request.Request(url, data=body, headers=headers, method=method)
    try:
        with urllib.request.urlopen(req) as r:
            ct = r.headers.get("Content-Type", "")
            raw_resp = r.read()
            if "json" in ct:
                return r.status, json.loads(raw_resp or "{}")
            return r.status, raw_resp
    except urllib.error.HTTPError as e:
        raw_resp = e.read()
        try:
            return e.code, json.loads(raw_resp or "{}")
        except json.JSONDecodeError:
            return e.code, raw_resp


def multipart(field, filename, content, content_type):
    boundary = "----testboundary123"
    parts = []
    parts.append(f"--{boundary}\r\n".encode())
    parts.append(
        f'Content-Disposition: form-data; name="{field}"; filename="{filename}"\r\n'.encode()
    )
    parts.append(f"Content-Type: {content_type}\r\n\r\n".encode())
    parts.append(content)
    parts.append(f"\r\n--{boundary}--\r\n".encode())
    body = b"".join(parts)
    return body, f"multipart/form-data; boundary={boundary}"


def multipart_form(fields, file_field, filename, content, content_type):
    """fields: {name: value} metin alanları + tek dosya alanı içeren multipart gövde."""
    boundary = "----testboundary456"
    parts = []
    for name, value in fields.items():
        parts.append(f"--{boundary}\r\n".encode())
        parts.append(f'Content-Disposition: form-data; name="{name}"\r\n\r\n'.encode())
        parts.append(f"{value}\r\n".encode())
    parts.append(f"--{boundary}\r\n".encode())
    parts.append(
        f'Content-Disposition: form-data; name="{file_field}"; filename="{filename}"\r\n'.encode()
    )
    parts.append(f"Content-Type: {content_type}\r\n\r\n".encode())
    parts.append(content)
    parts.append(f"\r\n--{boundary}--\r\n".encode())
    body = b"".join(parts)
    return body, f"multipart/form-data; boundary={boundary}"


def leaves(items, acc):
    for i in items:
        if i["children"]:
            leaves(i["children"], acc)
        else:
            acc.append(i)
    return acc


ok = 0
fail = 0


def check(name, cond):
    global ok, fail
    cond = bool(cond)
    print(("  PASS " if cond else "  FAIL ") + name)
    ok += cond
    fail += not cond


# ============================================================ Girişler
_, t = call("POST", "/auth/login", form={"username": "admin", "password": "admin123"})
admin = t["access_token"]
_, t = call("POST", "/auth/login", form={"username": "usta", "password": "usta123"})
usta = t["access_token"]
check("Giriş (admin + usta)", admin and usta)

s, _ = call("GET", "/users", token=usta)
check("Usta admin ekranına erişemez (403)", s == 403)

# ============================================================ Kullanıcı kaydı
s, r = call(
    "POST",
    "/auth/register",
    data={"username": "test.kayit", "full_name": "Test Kayıt", "password": "sifre123"},
)
check("Kayıt başarılı", s == 200)

s, r = call("POST", "/auth/login", form={"username": "test.kayit", "password": "sifre123"})
check("Onaysız kullanıcı giriş yapamaz (403)", s == 403)

_, users = call("GET", "/users", token=admin)
newu = next(u for u in users if u["username"] == "test.kayit")
check("Yeni kullanıcı pasif (onay bekliyor) oluştu", newu["is_active"] is False)
s, r = call("PATCH", f"/users/{newu['id']}", token=admin, data={"is_active": True})
check("Admin kullanıcıyı onayladı", s == 200 and r["is_active"] is True)

s, r = call("POST", "/auth/login", form={"username": "test.kayit", "password": "sifre123"})
check("Onaylı kullanıcı giriş yapabilir", s == 200 and "access_token" in r)

# ============================================================ Otomatik bakım açma
# 3801 -> 1.024.000 km: 1.025.000 bakımı aralığa girer (1.012.500 - 1.027.500)
_, vehicles = call("GET", "/vehicles", token=admin)
v = next(x for x in vehicles if x["code"] == "3801")
call("PUT", f"/vehicles/{v['id']}/km", token=admin, data={"current_km": 1_024_000})
_, ms = call("GET", "/maintenances", token=admin)
m = [x for x in ms if x["vehicle"]["id"] == v["id"] and x["target_km"] == 1_025_000]
check("Otomatik bakım açıldı (1.025.000 km, Bekliyor)", len(m) == 1 and m[0]["status"] == "pending")

call("PUT", f"/vehicles/{v['id']}/km", token=admin, data={"current_km": 1_025_500})
_, ms = call("GET", "/maintenances", token=admin)
m2 = [x for x in ms if x["vehicle"]["id"] == v["id"] and x["target_km"] == 1_025_000]
check("Aynı bakım ikinci kez açılmadı", len(m2) == 1)

mid = m[0]["id"]
_, detail = call("GET", f"/maintenances/{mid}", token=usta)
lv = leaves(detail["items"], [])
check(f"Gerçek bakım ağacı snapshot kopyalandı ({len(lv)} kontrol maddesi)", len(lv) > 20)

parent = detail["items"][0]
s, r = call("POST", f"/items/{parent['id']}/complete", token=usta)
check("Üst düğüm işaretlenemedi (400)", s == 400)

for leaf in lv:
    call("POST", f"/items/{leaf['id']}/complete", token=usta)
_, detail = call("GET", f"/maintenances/{mid}", token=admin)
check("Tüm üst düğümler otomatik tamamlandı",
      all(i["is_completed"] and i["is_locked"] for i in detail["items"]))
check("Bakım durumu: Tamamlandı", detail["status"] == "completed")
lv2 = leaves(detail["items"], [])
check("Kim yaptı kaydı tutuldu",
      all(x["completed_by"] and x["completed_by"]["username"] == "usta" and x["completed_at"] for x in lv2))

s, _ = call("POST", f"/items/{lv2[0]['id']}/complete", token=usta)
check("Kilitli madde değiştirilemedi (400)", s == 400)
s, _ = call("POST", f"/items/{lv2[0]['id']}/unlock", token=usta)
check("Usta kilit açamadı (403)", s == 403)

s, _ = call("POST", f"/items/{lv2[0]['id']}/unlock", token=admin)
_, detail = call("GET", f"/maintenances/{mid}", token=admin)
lv3 = leaves(detail["items"], [])
target = next(x for x in lv3 if x["id"] == lv2[0]["id"])
check("Admin kilidi kaldırdı, madde sıfırlandı", s == 200 and not target["is_completed"])
check("Bakım tekrar 'Devam Ediyor' oldu", detail["status"] == "in_progress")

call("POST", f"/items/{target['id']}/complete", token=usta)
_, detail = call("GET", f"/maintenances/{mid}", token=admin)
check("Yeniden tamamlanınca bakım tekrar bitti", detail["status"] == "completed")

# ============================================================ Plana özel bakım ağacı editörü
_, plans = call("GET", "/plans", token=admin)
plan25 = next(p for p in plans if p["interval_km"] == 25000)
plan50 = next(p for p in plans if p["interval_km"] == 50000)

s, r = call("GET", f"/tree?plan_id={plan25['id']}", token=admin)
check("25.000 km planının kendi ağacı var (ARAÇ ÜSTÜ/İÇİ/ALTI)",
      s == 200 and {n["name"] for n in r} == {"ARAÇ ÜSTÜ", "ARAÇ İÇİ", "ARAÇ ALTI"})

# ============================================================ Kademeli (kümülatif) bakım
# 150.000 km hem 25.000 hem 50.000'e tam bölünür -> tek iş emri, yönetici plan 50.000,
# checklist iki planın birleşimi (grup başlıkları altında).
_, kum_v = call("POST", "/vehicles", token=admin,
                data={"code": "KUM1", "name": "Kümülatif Test", "current_km": 150000})
_, kum_ms = call("GET", "/maintenances", token=admin)
kum_ms = [x for x in kum_ms if x["vehicle"]["id"] == kum_v["id"]]
kum_150 = next((x for x in kum_ms if x["target_km"] == 150000), None)
check("150.000 km'de tek iş emri açıldı (25k+50k birleşik)",
      kum_150 is not None and kum_150["plan"]["interval_km"] == 50000)

kum_detail = None
if kum_150:
    _, kum_detail = call("GET", f"/maintenances/{kum_150['id']}", token=admin)
if kum_detail:
    root_names = {n["name"] for n in kum_detail["items"]}
    check("Birleşik iş emrinde 25.000 ve 50.000 grupları var",
          "25.000 km Bakımı Kontrolleri" in root_names and "50.000 km Bakımı Kontrolleri" in root_names)
    kum_leaves = leaves(kum_detail["items"], [])
    check("Birleşik iş emrinin madde sayısı tek plandan fazla", len(kum_leaves) > len(lv))

    for leaf in kum_leaves:
        call("POST", f"/items/{leaf['id']}/complete", token=usta)
    _, kum_detail2 = call("GET", f"/maintenances/{kum_150['id']}", token=admin)
    check("Tüm gruplar bitince birleşik iş emri de tamamlandı", kum_detail2["status"] == "completed")

# 75.000 km sadece 25.000'e bölünür -> tek plan, sarmalayıcı grup YOK (düz ağaç)
_, kum_v2 = call("POST", "/vehicles", token=admin,
                 data={"code": "KUM2", "name": "Kümülatif Test 2", "current_km": 75000})
_, kum_ms2 = call("GET", "/maintenances", token=admin)
kum_75 = next((x for x in kum_ms2 if x["vehicle"]["id"] == kum_v2["id"] and x["target_km"] == 75000), None)
check("75.000 km'de tek plan (25.000) yönetici", kum_75 is not None and kum_75["plan"]["interval_km"] == 25000)
if kum_75:
    _, kum_detail3 = call("GET", f"/maintenances/{kum_75['id']}", token=admin)
    root_names3 = {n["name"] for n in kum_detail3["items"]}
    check("Tek plan durumunda sarmalayıcı grup yok (doğrudan ARAÇ ÜSTÜ/İÇİ/ALTI)",
          root_names3 == {"ARAÇ ÜSTÜ", "ARAÇ İÇİ", "ARAÇ ALTI"})

s, r2 = call("GET", f"/tree?plan_id={plan50['id']}", token=admin)
check("50.000 km planının ağacı 25.000'den farklı (bağımsız)", s == 200 and r2 != r)

# Araç İçi -> alt bölüm ekleme/çıkarma (sınırsız seviye)
arac_ici = next(n for n in r if n["name"] == "ARAÇ İÇİ")
s, yeni_bolum = call(
    "POST", "/tree", token=admin,
    data={"name": "Yolcu Kabini", "parent_id": arac_ici["id"], "plan_id": plan25["id"]},
)
check("Araç İçi altına yeni alt bölüm eklendi", s == 200 and yeni_bolum["name"] == "Yolcu Kabini")

s, alt_madde = call(
    "POST", "/tree", token=admin,
    data={"name": "Koltuk Kontrolü", "parent_id": yeni_bolum["id"]},
)
check("Yolcu Kabini altına kontrol maddesi eklendi", s == 200)

s, _ = call("DELETE", f"/tree/{alt_madde['id']}", token=admin)
check("Kontrol maddesi silindi", s == 200)
s, _ = call("DELETE", f"/tree/{yeni_bolum['id']}", token=admin)
check("Alt bölüm silindi", s == 200)

_, n1 = call("POST", "/tree", token=admin, data={"name": "Test Bölümü", "plan_id": plan25["id"]})
_, n2 = call("POST", "/tree", token=admin, data={"name": "Alt Madde", "parent_id": n1["id"]})
s, _ = call("PATCH", f"/tree/{n1['id']}", token=admin, data={"parent_id": n2["id"]})
check("Döngü koruması çalıştı (400)", s == 400)
s, _ = call("DELETE", f"/tree/{n1['id']}", token=admin)
check("Bölüm silindi (alt maddelerle birlikte)", s == 200)

# ============================================================ Dashboard
s, dash = call("GET", "/dashboard", token=admin)
check("Dashboard yüklendi", s == 200)
check("Dashboard bakım sayaçları tutarlı",
      dash["completed_count"] >= 1 and dash["vehicle_count"] >= 2)
check("Dashboard yaklaşan bakımlar listesi dolu", len(dash["upcoming"]) > 0)
check("Dashboard son tamamlananlar listesinde bizim bakım var",
      any(x["id"] == mid for x in dash["recent_completed"]))

# ============================================================ Araç Excel içe/dışa aktarma
import openpyxl

wb = openpyxl.Workbook()
ws = wb.active
ws.append(["Araç Kodu", "Araç Adı", "KM", "Üretici", "Model", "Şasi No", "Not"])
ws.append(["3801", "Tramvay 3801", 1_030_000, "Hyundai Rotem", "Sirio", "SIR-3801", "İçe aktarma testi"])
ws.append(["3899", "Test Aracı", 5000, "Hyundai Rotem", "Sirio", "SIR-3899", ""])
buf = io.BytesIO()
wb.save(buf)
xlsx_bytes = buf.getvalue()

body, ct = multipart("file", "araclar.xlsx", xlsx_bytes,
                      "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
s, r = call("POST", "/vehicles/import", token=admin, raw=body, raw_ct=ct)
check("Excel içe aktarma çalıştı (1 yeni, 1 güncelleme)",
      s == 200 and r["created"] == 1 and r["updated"] == 1)

_, vehicles = call("GET", "/vehicles", token=admin)
v01 = next(x for x in vehicles if x["code"] == "3801")
v99 = next((x for x in vehicles if x["code"] == "3899"), None)
check("İçe aktarılan KM güncellendi", v01["current_km"] == 1_030_000)
check("İçe aktarılan diğer bilgiler kaydedildi", v01["chassis_no"] == "SIR-3801")
check("Yeni araç oluşturuldu", v99 is not None and v99["current_km"] == 5000)
check("last_km_update_at dolduruldu", v01["last_km_update_at"] is not None)

s, exported = call("GET", "/vehicles/export", token=admin)
check("Excel dışa aktarma çalıştı", s == 200 and isinstance(exported, bytes) and len(exported) > 1000)

# ============================================================ 3D Model + Parça Eşleşmesi
s, flat = call("GET", f"/tree/flat?plan_id={plan25['id']}", token=admin)
check("Düz ağaç listesi (parça ataması için) döndü", s == 200 and len(flat) > 10)
flat_node_id = flat[0]["id"]

s, models3d = call("GET", "/models3d", token=admin)
check("3D model listesi döndü (seed'den en az 1 model beklenir)", s == 200)

glb_bytes = b"glTF" + (b"\x00" * 16)  # gerçek olmayan minik ikili içerik, yalnızca dosya akışı testi için
body, ct = multipart_form({"label": "Test Modeli"}, "file", "test.glb", glb_bytes, "model/gltf-binary")
s, model3d = call("POST", "/models3d", token=admin, raw=body, raw_ct=ct)
check("3D model yüklendi", s == 200 and bool(model3d.get("filename")))
model3d_id = model3d["id"]

glb_req = urllib.request.Request(f"http://localhost:8000/uploads/models/{model3d['filename']}")
try:
    with urllib.request.urlopen(glb_req) as gr:
        glb_status = gr.status
except urllib.error.HTTPError as e:
    glb_status = e.code
check("Yüklenen 3D model dosyası erişilebilir", glb_status == 200)

s, r = call("POST", f"/models3d/{model3d_id}/hotspots", token=admin,
            data={"plan_id": plan25["id"], "tree_node_id": flat_node_id,
                  "mesh_name": "Coupler Scharfenberg", "label": "Kuplör"})
check("Parça eşleşmesi (mesh->madde) oluşturuldu", s == 200)
part_hotspot_id = r["id"]

s, r = call("PATCH", f"/model-hotspots/{part_hotspot_id}", token=admin, data={"label": "Kuplör (güncel)"})
check("Parça eşleşmesi güncellendi", s == 200 and r["label"] == "Kuplör (güncel)")

s, _ = call("POST", f"/models3d/{model3d_id}/hotspots", token=usta,
            data={"plan_id": plan25["id"], "tree_node_id": flat_node_id, "mesh_name": "X"})
check("Usta parça eşleşmesi oluşturamaz (403)", s == 403)

s, listed = call("GET", f"/models3d/{model3d_id}/hotspots?plan_id={plan25['id']}", token=admin)
check("Plana özel parça eşleşmeleri listelendi", s == 200 and any(h["id"] == part_hotspot_id for h in listed))

s, m3d_status = call("GET", f"/maintenances/{mid}/model3d", token=admin)
check("Bakım detayında 3D model durumu döndü", s == 200)
match = next((d for d in m3d_status if d["id"] == model3d_id), None)
check("3D model bakım görünümünde eşleşti", match is not None)
if match:
    hs = next((x for x in match["hotspots"] if x["hotspot"]["id"] == part_hotspot_id), None)
    # 'mid' bakımı önceki adımlarda zaten tamamlanmıştı, o yüzden kök düğüme
    # bağlı parça da tamamlanmış (tikli) görünmeli.
    check("Parça durumu (tamamlandı) doğru döndü", hs is not None and hs["is_completed"] is True)

s, _ = call("DELETE", f"/model-hotspots/{part_hotspot_id}", token=admin)
check("Parça eşleşmesi silindi", s == 200)
s, _ = call("DELETE", f"/models3d/{model3d_id}", token=admin)
check("3D model silindi", s == 200)

glb_req2 = urllib.request.Request(f"http://localhost:8000/uploads/models/{model3d['filename']}")
try:
    urllib.request.urlopen(glb_req2)
    glb_status2 = 200
except urllib.error.HTTPError as e:
    glb_status2 = e.code
check("3D model silinince dosya da kaldırıldı", glb_status2 == 404)

# ============================================================ 400.000 km planı silindi
_, all_plans = call("GET", "/plans", token=admin)
check("400.000 km planı artık yok", not any(p["interval_km"] == 400000 for p in all_plans))

# ============================================================ Plan silme (DELETE /api/plans/{id})
s, del_test_plan = call(
    "POST", "/plans", token=admin,
    data={"name": "Silme Testi", "interval_km": 660000, "lower_pct": 10, "upper_pct": 10},
)
check("Silme testi için geçici plan oluşturuldu", s == 200)

s, _ = call("DELETE", f"/plans/{del_test_plan['id']}", token=usta)
check("Usta plan silemez (403)", s == 403)

s, r = call("DELETE", f"/plans/{del_test_plan['id']}", token=admin)
check("Geçmişi olmayan (bakımı hiç açılmamış) plan tamamen silindi", s == 200 and r["deactivated"] is False)
_, plans_after_del = call("GET", "/plans", token=admin)
check("Silinen plan artık listede yok", not any(p["id"] == del_test_plan["id"] for p in plans_after_del))

s, _ = call("DELETE", "/plans/999999", token=admin)
check("Var olmayan plan silinmeye çalışılınca 404", s == 404)

# plan25'in zaten açılmış bakımları var (yukarıdaki testlerden) -> silinmemeli, pasife alınmalı
s, r = call("DELETE", f"/plans/{plan25['id']}", token=admin)
check("Geçmişi olan plan (25.000 km) silinmedi, pasife alındı", s == 200 and r["deactivated"] is True)
_, plans_after_del2 = call("GET", "/plans", token=admin)
p25_after = next(p for p in plans_after_del2 if p["id"] == plan25["id"])
check("Plan pasif durumda görünüyor", p25_after["is_active"] is False)
# Sonraki testleri etkilememesi için tekrar aktifleştiriyoruz.
call("PATCH", f"/plans/{plan25['id']}", token=admin, data={"is_active": True})

# ============================================================ Bakım açma hatası düzeltmesi (500k/300k)
# Önceki (hatalı) algoritmada, 500.000'i biraz geçen bir araç için 600.000 gibi
# uzak bir hedef de aday oluyor ve onu "yöneten" 300.000 km'lik planın geniş
# toleransıyla yanlışlıkla açılabiliyordu. Düzeltilmiş algoritma yalnızca en
# yakın kontrol noktasını (burada tam 500.000) aday sayar.
_, base_v = call("POST", "/vehicles", token=admin,
                  data={"code": "BASE500K", "name": "Taban Birim Testi", "current_km": 500_100})
_, base_ms = call("GET", "/maintenances", token=admin)
base_ms = [x for x in base_ms if x["vehicle"]["id"] == base_v["id"]]
check("500.000 km'yi biraz geçince sadece 500.000 bakımı açıldı",
      any(x["target_km"] == 500000 and x["plan"]["interval_km"] == 500000 for x in base_ms))
check("300.000 km bakımı hatalı/erken şekilde açılmadı",
      not any(x["plan"]["interval_km"] == 300000 for x in base_ms))
check("Uzak hedef (600.000 gibi) hiç açılmadı",
      not any(x["target_km"] == 600000 for x in base_ms))

# ============================================================ Excel ile self-servis yeni bakım tablosu
wb2 = openpyxl.Workbook()
ws2 = wb2.active
ws2.title = "Sayfa1"
ws2.cell(row=1, column=1, value=777000)
ws2.cell(row=2, column=1, value="TEST KATEGORİ")
ws2.cell(row=3, column=1, value="TEST ALT KATEGORİ")
ws2.cell(row=4, column=1, value="TEST BİLEŞEN")
ws2.cell(row=5, column=1, value="Test kontrol maddesi 1")
ws2.cell(row=6, column=1, value="Test kontrol maddesi 2")
buf2 = io.BytesIO()
wb2.save(buf2)
tree_xlsx_bytes = buf2.getvalue()

body, ct = multipart("file", "yeni_bakim.xlsx", tree_xlsx_bytes,
                      "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
s, _ = call("POST", "/tree/import-excel", token=usta, raw=body, raw_ct=ct)
check("Usta Excel ile yeni tablo oluşturamaz (403)", s == 403)

s, r = call("POST", "/tree/import-excel", token=admin, raw=body, raw_ct=ct)
check("Excel ile yeni bakım tablosu (Bakım Ağacı sayfasından) oluşturuldu",
      s == 200 and r["nodes_created"] > 0)
new_plan_id = next((p["id"] for p in r["plans"] if p["interval_km"] == 777000), None)
check("Yeni plan (777.000 km) Excel'den oluşturuldu", new_plan_id is not None)

if new_plan_id:
    s, new_tree = call("GET", f"/tree?plan_id={new_plan_id}", token=admin)
    check("Yeni planın kendi bakım ağacı ayrıca görüntülenebiliyor",
          s == 200 and any(n["name"] == "TEST KATEGORİ" for n in new_tree))
    call("DELETE", f"/plans/{new_plan_id}", token=admin)

# ============================================================ Dashboard yaklaşan bakımlar penceresi (-%5/+%10)
# Auto-open toleransı (planın kendi lower_pct/upper_pct'i) ile dashboard'un
# sabit önizleme penceresini birbirine karıştırmamak için, 500.000 km
# planının toleransını geçici olarak sıfıra çekiyoruz; böylece yalnızca
# dashboard'un kendi -%5/+%10 penceresi devrede olur.
_, plans_now = call("GET", "/plans", token=admin)
plan500 = next(p for p in plans_now if p["interval_km"] == 500000)
call("PATCH", f"/plans/{plan500['id']}", token=admin, data={"lower_pct": 0, "upper_pct": 0})

_, dash_near_v = call("POST", "/vehicles", token=admin,
                       data={"code": "DASHNEAR", "name": "Dashboard Yakın", "current_km": 480_000})
_, dash_far_v = call("POST", "/vehicles", token=admin,
                      data={"code": "DASHFAR", "name": "Dashboard Uzak", "current_km": 260_000})

s, dash2 = call("GET", "/dashboard", token=admin)
near_shown = any(
    u["vehicle"]["id"] == dash_near_v["id"] and u["target_km"] == 500000 for u in dash2["upcoming"]
)
far_shown = any(u["vehicle"]["id"] == dash_far_v["id"] for u in dash2["upcoming"])
check("Hedefe yakın araç (480.000/500.000 km) yaklaşanlar listesinde", near_shown)
check("Hedefe uzak araç (260.000/500.000 km) yaklaşanlar listesinde DEĞİL", not far_shown)

# Aynı pencere mantığı Araçlar sayfasındaki /vehicles/alerts uç noktasında da
# (dashboard'un aksine TÜM araçlar için, 8 ile sınırlı olmadan) kullanılabilir olmalı.
s, alerts = call("GET", "/vehicles/alerts", token=admin)
check("Araç uyarıları uç noktası çalıştı", s == 200)
alert_near_shown = any(
    u["vehicle"]["id"] == dash_near_v["id"] and u["target_km"] == 500000 for u in alerts
)
check("Yakın araç /vehicles/alerts listesinde de var", alert_near_shown)

call("PATCH", f"/plans/{plan500['id']}", token=admin, data={"lower_pct": 50, "upper_pct": 10})

# ============================================================ Araç silme (DELETE /api/vehicles/{id})
s, del_v = call("POST", "/vehicles", token=admin,
                data={"code": "DELTEST", "name": "Silme Testi Aracı", "current_km": 1000})
check("Silme testi için geçici araç oluşturuldu", s == 200)

s, _ = call("DELETE", f"/vehicles/{del_v['id']}", token=usta)
check("Usta araç silemez (403)", s == 403)

s, r = call("DELETE", f"/vehicles/{del_v['id']}", token=admin)
check("Geçmişi olmayan (hiç bakımı açılmamış) araç tamamen silindi", s == 200)
_, vehicles_after_del = call("GET", "/vehicles", token=admin)
check("Silinen araç artık listede yok", not any(x["id"] == del_v["id"] for x in vehicles_after_del))

# kum_v'nin (yukarıda oluşturulan KUM1) zaten açılmış/tamamlanmış bakımı var
# -> silinmemeli, pasife alınmalı (geçmiş kaybolmasın).
s, r = call("DELETE", f"/vehicles/{kum_v['id']}", token=admin)
check("Geçmişi olan araç silinmedi, pasife alındı", s == 200)
_, vehicles_after_del2 = call("GET", "/vehicles", token=admin)
kum_v_after = next(x for x in vehicles_after_del2 if x["id"] == kum_v["id"])
check("Araç pasif durumda görünüyor", kum_v_after["is_active"] is False)

# ============================================================ Araç geçmişi görüntüleme (vehicle_id filtresi)
s, v_history = call("GET", f"/maintenances?vehicle_id={v['id']}", token=admin)
check("Araca özel bakım geçmişi filtrelenebiliyor", s == 200 and len(v_history) > 0)
check("Filtrelenen geçmişteki tüm kayıtlar doğru araca ait",
      all(x["vehicle"]["id"] == v["id"] for x in v_history))

print(f"\nSonuç: {ok} başarılı, {fail} hatalı")
raise SystemExit(1 if fail else 0)
