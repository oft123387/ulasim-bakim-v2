"""build_maintenances_workbook — Bakım Geçmişi Excel dışa aktarımında
periyoda (bakım planına) göre sayfa bölünmesi.

Kullanıcı isteği: "bakım geçmişi excel indirirken sheet sheet ayrı ayrı
böl mesela 200 bin bakımı için olanlar 500 bin bakımı için olanlar gibi."
Tek durum (ör. yalnızca Tamamlandı — Bakım Geçmişi'nin her zamanki hali)
içeren bir dışa aktarımda, farklı periyotlar (planlar) artık tek kalabalık
sayfada değil, her biri kendi sayfasında ayrı ayrı görünür. Tek periyot
varsa (tüm sonuçlar aynı plana ait) gereksiz bölünme olmaz."""

import io
from types import SimpleNamespace
from datetime import datetime

import openpyxl

from app.excel_utils import build_maintenances_workbook


def _plan(name, interval_km=0, interval_months=None, plan_type="km"):
    return SimpleNamespace(
        name=name, interval_km=interval_km, interval_months=interval_months, plan_type=plan_type,
    )


def _vehicle(code):
    return SimpleNamespace(code=code, name=f"Tramvay {code}")


def _maintenance(vehicle, plan, target_km, status="completed"):
    return SimpleNamespace(
        id=1,
        vehicle=vehicle,
        plan=plan,
        target_km=target_km,
        km_at_open=target_km,
        # Bakım yapılırken girilen güncel km (Faz — km sapma yüzdesi) — bu
        # testler dışa aktarım format/sayfa bölünme davranışını doğruluyor,
        # sapma hesaplamasını değil; None bırakmak compute_km_deviation_pct'in
        # deviation=None dönmesini sağlar (sütun boş görünür, hata vermez).
        completed_km=None,
        status=status,
        items=[],
        assigned_to=None,
        created_at=datetime.now(),
        started_at=datetime.now(),
        completed_at=datetime.now(),
    )


def test_single_status_multiple_plans_splits_into_sheets_per_period():
    plan_200k = _plan("200.000 km Bakımı", interval_km=200_000)
    plan_500k = _plan("500.000 km Bakımı", interval_km=500_000)
    rows = [
        _maintenance(_vehicle("3801"), plan_200k, 200_000),
        _maintenance(_vehicle("3802"), plan_200k, 400_000),
        _maintenance(_vehicle("3803"), plan_500k, 500_000),
    ]
    content = build_maintenances_workbook(rows, title="Bakım Geçmişi")
    wb = openpyxl.load_workbook(io.BytesIO(content))

    assert set(wb.sheetnames) == {"200.000 km Bakımı", "500.000 km Bakımı"}
    ws_200k = wb["200.000 km Bakımı"]
    # Başlık satırı + 2 veri satırı
    assert ws_200k.max_row == 3
    ws_500k = wb["500.000 km Bakımı"]
    assert ws_500k.max_row == 2


def test_single_status_single_plan_keeps_one_sheet():
    plan_200k = _plan("200.000 km Bakımı", interval_km=200_000)
    rows = [
        _maintenance(_vehicle("3801"), plan_200k, 200_000),
        _maintenance(_vehicle("3802"), plan_200k, 200_000),
    ]
    content = build_maintenances_workbook(rows, title="Bakım Geçmişi")
    wb = openpyxl.load_workbook(io.BytesIO(content))

    assert wb.sheetnames == ["Bakım Geçmişi"]


def test_multiple_statuses_still_split_by_status_not_plan():
    plan_200k = _plan("200.000 km Bakımı", interval_km=200_000)
    plan_500k = _plan("500.000 km Bakımı", interval_km=500_000)
    rows = [
        _maintenance(_vehicle("3801"), plan_200k, 200_000, status="completed"),
        _maintenance(_vehicle("3802"), plan_500k, 500_000, status="pending"),
    ]
    content = build_maintenances_workbook(rows, title="Bakımlar")
    wb = openpyxl.load_workbook(io.BytesIO(content))

    assert set(wb.sheetnames) == {"Tamamlandı", "Bekliyor"}
