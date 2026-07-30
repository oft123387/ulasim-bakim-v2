"""GET /api/maintenances/{id}/export-items — bir iş emrinin kontrol
maddelerini kimin/ne zaman tamamladığını (ve varsa eklediği notu) Excel
olarak dışa aktaran uç.

Yalnızca en alt seviye (yaprak) maddeler satır olur; üst başlıklar "Yol"
sütununda breadcrumb olarak görünür, kendileri ayrı satır olmaz.
"""

import io

import openpyxl

from .conftest import auth


def test_export_maintenance_items_xlsx_contains_completer_and_timestamp(client, admin_token):
    admin = admin_token
    model = "ExportItemsTest Model"

    r = client.post(
        "/api/vehicles",
        headers=auth(admin),
        json={"code": "9601", "name": "Export Test Tramvay", "current_km": 47_000, "model": model},
    )
    assert r.status_code == 200, r.text
    vehicle = r.json()

    r = client.post(
        "/api/plans",
        headers=auth(admin),
        json={"name": f"{model} Planı", "plan_type": "km", "interval_km": 47_000, "vehicle_model": model},
    )
    assert r.status_code == 200, r.text
    plan = r.json()

    r = client.post(
        "/api/tree", headers=auth(admin),
        json={"name": "ARAÇ İÇİ", "plan_id": plan["id"]},
    )
    assert r.status_code == 200, r.text
    root = r.json()
    r = client.post(
        "/api/tree", headers=auth(admin),
        json={"name": "Fren balatası kontrolü", "parent_id": root["id"], "card_code": "EXP-01"},
    )
    assert r.status_code == 200, r.text

    r = client.post(
        "/api/maintenances/manual",
        headers=auth(admin),
        json={"vehicle_id": vehicle["id"], "plan_id": plan["id"], "target_km": 47_000},
    )
    assert r.status_code == 200, r.text
    maint = r.json()

    detail = client.get(f"/api/maintenances/{maint['id']}", headers=auth(admin)).json()
    leaf = detail["items"][0]["children"][0]
    assert leaf["name"] == "Fren balatası kontrolü"

    r = client.post(f"/api/items/{leaf['id']}/complete", headers=auth(admin))
    assert r.status_code == 200, r.text

    r = client.get(f"/api/maintenances/{maint['id']}/export-items?format=xlsx", headers=auth(admin))
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "attachment" in r.headers["content-disposition"]

    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    ws = wb.active
    header = [c.value for c in ws[1]]
    assert header == [
        "İş Emri", "Araç", "Periyot", "Yol", "Madde", "Kart Kodu",
        "Durum", "Tamamlayan", "Tamamlanma Tarihi/Saati",
        "Bulgu/Not", "Notu Ekleyen", "Not Tarihi/Saati",
    ]

    data_row = [c.value for c in ws[2]]
    assert data_row[0] == str(maint["id"])
    assert data_row[1] == "9601"
    assert data_row[4] == "Fren balatası kontrolü"
    assert data_row[5] == "EXP-01"
    assert data_row[6] == "Tamamlandı"
    assert data_row[7] == "Sistem Yöneticisi"  # admin kullanıcısının full_name'i (bkz. conftest seed)
    assert data_row[8]  # tamamlanma tarihi/saati dolu olmalı


def test_export_maintenance_items_requires_valid_maintenance(client, admin_token):
    r = client.get("/api/maintenances/999999/export-items", headers=auth(admin_token))
    assert r.status_code == 404


def test_export_maintenance_items_rejects_bad_format(client, admin_token):
    r = client.get("/api/maintenances/1/export-items?format=zzz", headers=auth(admin_token))
    assert r.status_code in (400, 404)
